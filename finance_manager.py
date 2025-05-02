# Operações financeiras
import pandas as pd
import json
import os
from datetime import datetime
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, NamedStyle

data_file = "data/transacoes.json"

class FinanceManager: # responsável por toda a lógica do controle financeiro
    def __init__(self):
        os.makedirs("data", exist_ok=True) # Cria a pasta data se ela não existir
        if not os.path.exists(data_file):
            with open(data_file, "w", encoding="utf-8") as f:
                json.dump([], f)
        self.transacoes = self.carregar_transacoes()
        """# DEBUG: Verifica se as transações foram carregadas corretamente
        print(f"Transações carregadas: {len(self.transacoes)}")
        for t in self.transacoes:
            print(t)"""

    def carregar_transacoes(self): # responsável por carregar as transações salvas no arquivo JSON
        with open(data_file, "r", encoding="utf-8") as f:
            try:
                return json.load(f) # json.load(f) converte o JSON em uma lista de dicionários Python
            except json.JSONDecodeError:
                return [] # retorna a lista vazia se o arquivo estiver vazio ou mal formatado
        
    def salvar_transacoes(self, transacoes): # responsável por salvar a lista de transações de volta ao arquivo JSON
        with open(data_file, "w") as f:
            json.dump(transacoes, f, indent=4) # indent=4 formata o JSON com identação de 4 espaços
            

    def adicionar_transacao(self, data, tipo, valor, categoria, forma_pagamento, descricao=""):
        transacoes = self.carregar_transacoes()
        nova = {
            "data": data,
            "tipo": tipo,  # "receita" ou "despesa"
            "valor": valor,
            "categoria": categoria,
            "forma_pagamento": forma_pagamento,
            "descricao": descricao
            }
        transacoes.append(nova)
        self.salvar_transacoes(transacoes)
        print("Transação adicionada com sucesso.")

    def exportar_para_excel(self, nome_arquivo="transacoes_exportadas.xlsx", data_inicio=None, data_fim=None):
        transacoes = self.carregar_transacoes()

        # Se datas forem fornecidas, filtra a lista
        if data_inicio and data_fim:
            data_i = datetime.strptime(data_inicio, "%d/%m/%Y")
            data_f = datetime.strptime(data_fim, "%d/%m/%Y")

            transacoes = [
                t for t in transacoes
                if data_i <= datetime.strptime(t["data"], "%d/%m/%Y") <= data_f
            ]

        if not transacoes:
            print("⚠️ Nenhuma transação encontrada no intervalo informado.")
            return
        
        df = pd.DataFrame(transacoes) # Converte a lista de dicionário em DataFrame
        nomes_colunas = {
            "tipo": "Tipo",
            "valor": "Valor",
            "categoria": "Categoria",
            "descricao": "Descrição",
            "data": "Data"
        }
        df.rename(columns=nomes_colunas, inplace=True)
        # Exportar para excel
        with pd.ExcelWriter(nome_arquivo, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Transações", index=False)
            # Acessa a planilha
            ws = writer.book["Transações"]
            # Cria um estilo de data no formato DD/MM/AAAA
            estilo_data = NamedStyle(name="data_brasileira", number_format="DD/MM/YYYY")
            # Descobre qual a coluna é "data"
            data_col_index = None
            for idx, cell in enumerate(ws[1], 1):
                if cell.value.lower() == "data":
                    data_col_index = idx
                    break
            # Aplica o estilo de data nas células da coluna correspondente
            if data_col_index:
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                    cell = row[data_col_index - 1]
                    try:
                        if isinstance(cell.value, datetime):
                            cell.style = estilo_data
                    except:
                        pass
            # congela cabeçalho
            ws.freeze_panes = "A2"
            # ativa filtros automáticos
            ws.auto_filter.ref = ws.dimensions
            # aplica estilos nas células do cabeçalho
            header_fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
            header_font = Font(bold=True)

            for col_num, cell in enumerate(ws[1], 1):
                cell.fill = header_fill
                cell.font = header_font
                # ajusta largura da coluna conforme conteúdo
                col_letter = get_column_letter(col_num)
                max_length = max(len(str(cell.value)) for cell in ws[col_letter])
                ws.column_dimensions[col_letter].width = max_length + 2
            
            # Estilizar entrada e saída
            entrada_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            saida_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            # identifica o índice da coluna "tipo"
            tipo_col_index = None
            for idx, cell in enumerate(ws[1], 1):
                if cell.value.lower() == "tipo":
                    tipo_col_index = idx
                    break

            # aplica cor as linhas com base no tipo de transação
            if tipo_col_index:
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row): # ignora cabeçalho
                    tipo = row[tipo_col_index - 1].value.lower()
                    if tipo == "receita":
                        for cell in row:
                            cell.fill = entrada_fill
                    elif tipo == "despesa":
                        for cell in row:
                            cell.fill = saida_fill

    def filtrar_por_data(self, data_inicio, data_fim):
        data_i = datetime.strptime(data_inicio, "%d/%m/%Y")
        data_f = datetime.strptime(data_fim, "%d/%m/%Y")
        # Converte string para datetime se necessário
        resultado = []
        for t in self.transacoes:
            try:
                data_t = datetime.strptime(t["data"], "%d/%m/%Y")
                if data_i <= data_t <= data_f:
                    resultado.append(t)
            except ValueError:
                continue
        return resultado
    
    def configurar_vencimento_fatura(self, dia_vencimento):
        with open("config.json", "w") as f:
            json.dump({"vencimento_fatura": str(dia_vencimento)}, f)

    def obter_vencimento_fatura(self):
        try:
            with open("config.json", "r") as f:
                config = json.load(f)
                return int(config.get("vencimento_fatura", 10))  # padrão: dia 10
        except:
            return 10       